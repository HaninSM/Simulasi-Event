#!/usr/bin/env python3
"""
Generate 5 Distinct Theme Variants of the Complete Mini-SIEM Presentation (.pptx)
All 5 presentations contain the EXACT SAME comprehensive 16-slide content.
Themes are inspired by high-end design systems (PPT Master: Pritzker 2026, Dujiangyan Heritage, Nordic Minimalist, Executive Boardroom, Modern Studio Clean).
All variants feature light, professional, executive backgrounds strictly avoiding geeky/terminal/IT tropes.
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

class PresentationTheme:
    def __init__(
        self,
        theme_id,
        theme_name,
        output_file,
        title_font,
        body_font,
        code_font,
        bg_color,
        card_bg,
        card_alt,
        card_border,
        text_primary,
        text_secondary,
        accent_primary,
        accent_secondary,
        color_danger,
        color_warning,
        color_success,
        color_info,
        has_top_bar=True,
        has_accent_line=False,
        round_card_corner=True,
        border_width=1.0,
        tag_style="pill" # "pill", "bracket", "rect"
    ):
        self.theme_id = theme_id
        self.theme_name = theme_name
        self.output_file = output_file
        self.title_font = title_font
        self.body_font = body_font
        self.code_font = code_font
        self.bg_color = bg_color
        self.card_bg = card_bg
        self.card_alt = card_alt
        self.card_border = card_border
        self.text_primary = text_primary
        self.text_secondary = text_secondary
        self.accent_primary = accent_primary
        self.accent_secondary = accent_secondary
        self.color_danger = color_danger
        self.color_warning = color_warning
        self.color_success = color_success
        self.color_info = color_info
        self.has_top_bar = has_top_bar
        self.has_accent_line = has_accent_line
        self.round_card_corner = round_card_corner
        self.border_width = border_width
        self.tag_style = tag_style

def get_themes():
    # 1. Pritzker Editorial (Architecture Monograph / Magazine Style)
    t1 = PresentationTheme(
        theme_id="pritzker",
        theme_name="Theme 1: Pritzker Editorial (Architecture Monograph)",
        output_file="variant1_pritzker_editorial.pptx",
        title_font="Georgia",
        body_font="Calibri",
        code_font="Consolas",
        bg_color=RGBColor(252, 251, 249),         # #fcfbf9 Warm Gallery Paper
        card_bg=RGBColor(255, 255, 255),          # #ffffff Pure White
        card_alt=RGBColor(245, 244, 240),         # #f5f4f0 Warm Stone Wash
        card_border=RGBColor(231, 229, 228),      # #e7e5e4 Fine Warm Stone Border
        text_primary=RGBColor(28, 25, 23),        # #1c1917 Deep Espresso Charcoal
        text_secondary=RGBColor(120, 113, 108),   # #78716c Umber Muted Slate
        accent_primary=RGBColor(180, 83, 9),      # #b45309 Warm Architectural Bronze
        accent_secondary=RGBColor(217, 119, 6),   # #d97706 Warm Ochre
        color_danger=RGBColor(153, 27, 27),       # #991b1b Terracotta Red
        color_warning=RGBColor(194, 65, 12),      # #c2410c Burnt Sienna
        color_success=RGBColor(21, 128, 61),      # #15803d Moss Olive Green
        color_info=RGBColor(180, 83, 9),          # #b45309 Architectural Bronze
        has_top_bar=False,
        has_accent_line=True,
        round_card_corner=False,
        border_width=0.8,
        tag_style="bracket"
    )

    # 2. Dujiangyan Heritage (Scholarly Engineering & Water Conservancy)
    t2 = PresentationTheme(
        theme_id="dujiangyan",
        theme_name="Theme 2: Dujiangyan Heritage (Scholarly Engineering)",
        output_file="variant2_dujiangyan_heritage.pptx",
        title_font="Palatino Linotype",
        body_font="Segoe UI",
        code_font="Consolas",
        bg_color=RGBColor(251, 251, 248),         # #fbfbf8 Soft Ivory Parchment
        card_bg=RGBColor(255, 255, 255),          # #ffffff Crisp Card
        card_alt=RGBColor(244, 247, 246),         # #f4f7f6 Soft Sage Wash
        card_border=RGBColor(209, 213, 219),      # #d1d5db River Pebble Border
        text_primary=RGBColor(31, 41, 55),        # #1f2937 Mineral Slate
        text_secondary=RGBColor(75, 85, 99),      # #4b5563 River Stone Muted
        accent_primary=RGBColor(6, 95, 70),       # #065f46 Deep River Jade / Pine
        accent_secondary=RGBColor(4, 120, 87),    # #047857 Mountain Emerald
        color_danger=RGBColor(185, 28, 28),       # #b91c1c Antique Vermilion
        color_warning=RGBColor(180, 83, 9),       # #b45309 Antique Brass
        color_success=RGBColor(5, 150, 105),      # #059669 Vibrant Jade
        color_info=RGBColor(6, 95, 70),           # #065f46 Deep Jade
        has_top_bar=True,
        has_accent_line=False,
        round_card_corner=True,
        border_width=1.0,
        tag_style="pill"
    )

    # 3. Nordic Slate Minimalist (Swiss International Typographic Style)
    t3 = PresentationTheme(
        theme_id="nordic",
        theme_name="Theme 3: Nordic Slate Minimalist (Swiss Style)",
        output_file="variant3_nordic_minimalist.pptx",
        title_font="Segoe UI Semibold",
        body_font="Segoe UI",
        code_font="Consolas",
        bg_color=RGBColor(255, 255, 255),         # #ffffff Pure Clean White
        card_bg=RGBColor(248, 250, 252),          # #f8fafc Soft Glacier Gray
        card_alt=RGBColor(241, 245, 249),         # #f1f5f9
        card_border=RGBColor(226, 232, 240),      # #e2e8f0 Hairline Arctic Slate
        text_primary=RGBColor(15, 23, 42),        # #0f172a Polar Graphite
        text_secondary=RGBColor(100, 116, 139),   # #64748b Cool Steel Muted
        accent_primary=RGBColor(2, 132, 199),     # #0284c7 Arctic Cobalt
        accent_secondary=RGBColor(30, 41, 59),    # #1e293b Deep Nordic Navy
        color_danger=RGBColor(220, 38, 38),       # #dc2626 Minimalist Coral Red
        color_warning=RGBColor(217, 119, 6),      # #d97706 Amber Alert
        color_success=RGBColor(16, 185, 129),     # #10b981 Mint Glacier
        color_info=RGBColor(2, 132, 199),         # #0284c7 Arctic Cobalt
        has_top_bar=True,
        has_accent_line=False,
        round_card_corner=False,
        border_width=1.0,
        tag_style="rect"
    )

    # 4. Executive Boardroom & Fintech (Institutional Wall Street Risk)
    t4 = PresentationTheme(
        theme_id="boardroom",
        theme_name="Theme 4: Executive Boardroom & Fintech",
        output_file="variant4_executive_boardroom.pptx",
        title_font="Trebuchet MS",
        body_font="Calibri",
        code_font="Consolas",
        bg_color=RGBColor(248, 250, 252),         # #f8fafc Oxford Off-White
        card_bg=RGBColor(255, 255, 255),          # #ffffff Pure White Card
        card_alt=RGBColor(241, 245, 249),         # #f1f5f9 Subtle Tint
        card_border=RGBColor(203, 213, 225),      # #cbd5e1 Banker's Silver Border
        text_primary=RGBColor(10, 25, 47),        # #0a192f Institutional Midnight Navy
        text_secondary=RGBColor(71, 85, 105),     # #475569 Banker's Slate Muted
        accent_primary=RGBColor(30, 58, 138),     # #1e3a8a Investment Royal Blue
        accent_secondary=RGBColor(180, 83, 9),    # #b45309 Sovereign Gold
        color_danger=RGBColor(153, 27, 27),       # #991b1b Imperial Crimson
        color_warning=RGBColor(217, 119, 6),      # #d97706 Treasury Amber
        color_success=RGBColor(4, 120, 87),       # #047857 Institutional Forest Green
        color_info=RGBColor(30, 58, 138),         # #1e3a8a Royal Blue
        has_top_bar=True,
        has_accent_line=True,
        round_card_corner=True,
        border_width=1.2,
        tag_style="pill"
    )

    # 5. Modern Studio Clean (Linear / Stripe Product Leadership)
    t5 = PresentationTheme(
        theme_id="studio",
        theme_name="Theme 5: Modern Studio Clean (Linear / Stripe UI Style)",
        output_file="variant5_modern_studio.pptx",
        title_font="Segoe UI",
        body_font="Segoe UI",
        code_font="Consolas",
        bg_color=RGBColor(255, 255, 255),         # #ffffff Modern Bright Canvas
        card_bg=RGBColor(248, 249, 254),          # #f8f9fe Soft Iris Tint
        card_alt=RGBColor(243, 244, 250),         # #f3f4fa
        card_border=RGBColor(224, 231, 255),      # #e0e7ff Modern Iris Hairline
        text_primary=RGBColor(30, 27, 75),        # #1e1b4b Ink Slate
        text_secondary=RGBColor(100, 116, 139),   # #64748b Modern Slate Muted
        accent_primary=RGBColor(79, 70, 229),     # #4f46e5 Studio Iris / Indigo
        accent_secondary=RGBColor(124, 58, 237),  # #7c3aed Studio Violet
        color_danger=RGBColor(225, 29, 72),       # #e11d48 Studio Ruby
        color_warning=RGBColor(217, 119, 6),      # #d97706 Modern Amber
        color_success=RGBColor(16, 185, 129),     # #10b981 Studio Mint Emerald
        color_info=RGBColor(79, 70, 229),         # #4f46e5 Studio Iris
        has_top_bar=True,
        has_accent_line=False,
        round_card_corner=True,
        border_width=1.0,
        tag_style="pill"
    )

    return [t1, t2, t3, t4, t5]

def add_theme_header(slide, theme, title_text, category_tag="MINI-SIEM COMMAND CENTER", subtitle_text=""):
    # 1. Canvas Background
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = theme.bg_color
    bg.line.fill.background()

    # 2. Top Decorative Bar
    if theme.has_top_bar:
        top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.08))
        top_bar.fill.solid()
        top_bar.fill.fore_color.rgb = theme.accent_primary
        top_bar.line.fill.background()

    # 3. Category Tag
    if theme.tag_style == "bracket":
        # Editorial bracket style: [ CATEGORY / PRD ]
        tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.35), Inches(8), Inches(0.32))
        tf = tag_box.text_frame
        p = tf.paragraphs[0]
        p.text = f"[  {category_tag.upper()}  ]"
        p.font.size = Pt(9)
        p.font.bold = True
        p.font.color.rgb = theme.accent_primary
        p.font.name = theme.title_font
    else:
        shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if theme.tag_style == "pill" else MSO_SHAPE.RECTANGLE
        pill = slide.shapes.add_shape(shape_type, Inches(0.8), Inches(0.35), Inches(4.2), Inches(0.32))
        pill.fill.solid()
        pill.fill.fore_color.rgb = theme.card_alt
        pill.line.color.rgb = theme.card_border
        pill.line.width = Pt(theme.border_width)

        tf = pill.text_frame
        p = tf.paragraphs[0]
        p.text = category_tag.upper()
        p.font.size = Pt(8.5)
        p.font.bold = True
        p.font.color.rgb = theme.accent_primary
        p.font.name = theme.body_font
        p.alignment = PP_ALIGN.CENTER

    # 4. Slide Title & Subtitle Box
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.68), Inches(11.7), Inches(0.8))
    tf2 = title_box.text_frame
    tf2.word_wrap = True
    p2 = tf2.paragraphs[0]
    p2.text = title_text
    p2.font.size = Pt(20)
    p2.font.bold = True
    p2.font.color.rgb = theme.text_primary
    p2.font.name = theme.title_font

    if subtitle_text:
        p_sub = tf2.add_paragraph()
        p_sub.text = subtitle_text
        p_sub.font.size = Pt(10.5)
        p_sub.font.color.rgb = theme.text_secondary
        p_sub.font.name = theme.body_font
        p_sub.space_before = Pt(3)

    # 5. Editorial Line Under Title (Optional)
    if theme.has_accent_line:
        line_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.5), Inches(11.733), Inches(0.015))
        line_shape.fill.solid()
        line_shape.fill.fore_color.rgb = theme.card_border
        line_shape.line.fill.background()

    # 6. Footer Line & Caption
    footer_line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(6.98), Inches(11.733), Inches(0.015))
    footer_line.fill.solid()
    footer_line.fill.fore_color.rgb = theme.card_border
    footer_line.line.fill.background()

    footer = slide.shapes.add_textbox(Inches(0.8), Inches(7.02), Inches(11.733), Inches(0.35))
    tff = footer.text_frame
    pf = tff.paragraphs[0]
    pf.text = f"Mini-SIEM Security Operations Platform  |  Tim 'JadiBegini'  |  {theme.theme_name}"
    pf.font.size = Pt(8.5)
    pf.font.color.rgb = theme.text_secondary
    pf.font.name = theme.body_font

def add_card_box(slide, theme, left, top, width, height, border_color=None, bg_color=None, side_stripe_color=None):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if theme.round_card_corner else MSO_SHAPE.RECTANGLE
    card = slide.shapes.add_shape(shape_type, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color or theme.card_bg
    card.line.color.rgb = border_color or theme.card_border
    card.line.width = Pt(theme.border_width)

    if side_stripe_color:
        stripe = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, Inches(0.12), height)
        stripe.fill.solid()
        stripe.fill.fore_color.rgb = side_stripe_color
        stripe.line.fill.background()

    return card

def build_variant_deck(theme):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5) # 16:9 Widescreen
    blank_layout = prs.slide_layouts[6]

    # ==================== SLIDE 1: COVER SLIDE ====================
    s1 = prs.slides.add_slide(blank_layout)
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = theme.bg_color
    bg1.line.fill.background()

    # Main Card
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if theme.round_card_corner else MSO_SHAPE.RECTANGLE
    card_main = s1.shapes.add_shape(shape_type, Inches(1.0), Inches(1.0), Inches(11.333), Inches(5.5))
    card_main.fill.solid()
    card_main.fill.fore_color.rgb = theme.card_bg
    card_main.line.color.rgb = theme.card_border
    card_main.line.width = Pt(theme.border_width * 1.5)

    # Left Accent Bar
    left_bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(1.0), Inches(0.18), Inches(5.5))
    left_bar.fill.solid()
    left_bar.fill.fore_color.rgb = theme.accent_primary
    left_bar.line.fill.background()

    tb = s1.shapes.add_textbox(Inches(1.8), Inches(1.5), Inches(9.8), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True

    p_badge = tf.paragraphs[0]
    p_badge.text = f"ENTERPRISE CYBERSECURITY SIMULATION  •  {theme.theme_name.upper()}"
    p_badge.font.size = Pt(10.5)
    p_badge.font.bold = True
    p_badge.font.color.rgb = theme.accent_primary
    p_badge.font.name = theme.body_font

    p_title = tf.add_paragraph()
    p_title.text = "Mini-SIEM Command Center"
    p_title.font.size = Pt(32)
    p_title.font.bold = True
    p_title.font.color.rgb = theme.text_primary
    p_title.font.name = theme.title_font
    p_title.space_before = Pt(8)

    p_sub = tf.add_paragraph()
    p_sub.text = "Platform Telemetri Heterogen Real-Time, Gatekeeper RFC Validation, SOAR Active Defense, Human-Supervised AI Adaptation & Korelasi 9 Skenario Insiden Siber/OT"
    p_sub.font.size = Pt(13)
    p_sub.font.color.rgb = theme.text_secondary
    p_sub.font.name = theme.body_font
    p_sub.space_before = Pt(14)

    p_meta = tf.add_paragraph()
    p_meta.text = (
        "Dokumen Induk: mini_siem_product_requirements_document.md (V2.2)\n"
        "Tim Pengembang: JadiBegini (Bagus Hanindyo Kusuma, Anand Ferdiansyah, Fitus Yayan Ofnaio Zebua)\n"
        "Framework & Stack: Python 3.11, FastAPI, Pydantic V2, SOAR Active Quarantine, MQTT Broker, NIST SP 800-92"
    )
    p_meta.font.size = Pt(10.5)
    p_meta.font.color.rgb = theme.accent_secondary
    p_meta.font.name = theme.body_font
    p_meta.space_before = Pt(20)

    # ==================== SLIDE 2: PROFIL TIM & FILOSOFI SIEM ====================
    s2 = prs.slides.add_slide(blank_layout)
    add_theme_header(s2, theme, "Profil Tim 'JadiBegini' & Filosofi Fundamental SIEM", "01 / TEAM & PHILOSOPHY (PRD §1.2 - 1.3)")

    team_members = [
        ("Bagus Hanindyo Kusuma", "TEAM LEAD / SCRUM MASTER", "Scrum Master & naturally agile coach dengan gaya kepemimpinan servant leadership. Berpengalaman luas sebagai Program Manager dan IT Project Manager dalam tata kelola siklus pengembangan sistem.", theme.accent_primary),
        ("Anand Ferdiansyah", "SECURITY RESEARCHER", "Mahasiswa Aktif Program Karyawan di Universitas Pamulang. Berfokus pada riset validasi skema log heterogen, gatekeeper ingestion, dan mitigasi time drift pada sistem keamanan siber.", theme.color_info),
        ("Fitus Yayan Ofnaio Zebua", "OPERATIONS ANALYST", "Mahasiswa Aktif Program Karyawan di Universitas Pamulang dan aktif berdedikasi sebagai Driver Grab. Berfokus pada analisis operasional korelasi insiden dan simulasi skenario lapangan.", theme.accent_secondary)
    ]

    for idx, (tname, trole, tbio, tcolor) in enumerate(team_members):
        x = Inches(0.8 + idx * 3.95)
        add_card_box(s2, theme, x, Inches(1.6), Inches(3.75), Inches(2.3), side_stripe_color=tcolor)

        tb = s2.shapes.add_textbox(x + Inches(0.25), Inches(1.75), Inches(3.35), Inches(2.0))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = tname
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = theme.text_primary
        p.font.name = theme.title_font

        p_r = tf.add_paragraph()
        p_r.text = trole
        p_r.font.size = Pt(9)
        p_r.font.bold = True
        p_r.font.color.rgb = tcolor
        p_r.font.name = theme.body_font
        p_r.space_before = Pt(2)

        p_b = tf.add_paragraph()
        p_b.text = tbio
        p_b.font.size = Pt(9.5)
        p_b.font.color.rgb = theme.text_secondary
        p_b.font.name = theme.body_font
        p_b.space_before = Pt(6)

    phils = [
        ("🏢 Analogi Gedung Perkantoran Raksasa", "SIEM diibaratkan ruang kendali terpadu gedung bertingkat ratusan pintu & CCTV. Kartu akses 'Budi' di-tap di turnstile lobi sementara workstation Budi aktif di lantai 15 memicu alarm korelasi seketika.", theme.accent_secondary),
        ("🛡️ SIEM vs Antivirus & 'SI' vs 'EM'", "Antivirus hanya menjaga 1 endpoint lokal; SIEM mengawasi seluruh perimeter. 'Security Information' mengarsipkan data kepatuhan; 'Event Management' mengorelasikan log secara real-time saat insiden berlangsung.", theme.color_success)
    ]

    for idx, (ptitle, pdesc, pcolor) in enumerate(phils):
        x = Inches(0.8 + idx * 5.95)
        add_card_box(s2, theme, x, Inches(4.1), Inches(5.75), Inches(2.55), side_stripe_color=pcolor)

        tb = s2.shapes.add_textbox(x + Inches(0.28), Inches(4.25), Inches(5.3), Inches(2.25))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = ptitle
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = theme.text_primary
        p.font.name = theme.title_font

        p_d = tf.add_paragraph()
        p_d.text = pdesc
        p_d.font.size = Pt(10.5)
        p_d.font.color.rgb = theme.text_secondary
        p_d.font.name = theme.body_font
        p_d.space_before = Pt(6)

    # ==================== SLIDE 3: RINGKASAN PRODUK ====================
    s3 = prs.slides.add_slide(blank_layout)
    add_theme_header(s3, theme, "Ringkasan Produk & 6 Nilai Strategis Mini-SIEM", "02 / PRODUCT OVERVIEW (PRD §1.1)")

    cards = [
        ("1. Gatekeeper Ingestion", "Normalisasi kontrak JSON kanonikal dari beragam sumber log mentah. Menolak data cacat dan clock drift via HTTP 422.", theme.accent_primary),
        ("2. Rule Correlation Engine", "Mesin korelasi sliding window stateful di memori mendeteksi serangan L7, brute force, insider threat, dan anomali jam operasional.", theme.color_info),
        ("3. SOAR Active Defense", "Respons insiden aktif memutus serangan di gerbang: Host Quarantine & User Suspension menolak paket penyerang via HTTP 403.", theme.color_success),
        ("4. Hot vs Cold Storage (APT)", "Mengatasi temporal blind spot serangan berjarak puluhan hari menggunakan Persistent Stateful Greylist dan rehidrasi forensik.", theme.accent_secondary),
        ("5. Universal Log Adapter", "Membaca 27 berkas log fisik autentik (Nginx, Syslog RFC 3164, CSV, PGAudit JSONL, MQTT JSONL) dengan timestamp shifting dinamis.", theme.color_warning),
        ("6. OT Hardware Monitoring", "Memonitor telemetri sensor dan kontrol hardware fasilitas datacenter via protokol MQTT untuk mencegah sabotase thermal katastropik.", theme.color_danger)
    ]

    for idx, (ctitle, cdesc, ccolor) in enumerate(cards):
        col = idx % 3
        row = idx // 3
        x = Inches(0.8 + col * 3.95)
        y = Inches(1.6 + row * 2.5)

        add_card_box(s3, theme, x, y, Inches(3.75), Inches(2.3), side_stripe_color=ccolor)

        tb = s3.shapes.add_textbox(x + Inches(0.25), y + Inches(0.2), Inches(3.3), Inches(1.9))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = ctitle
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = theme.text_primary
        p.font.name = theme.title_font

        p2 = tf.add_paragraph()
        p2.text = cdesc
        p2.font.size = Pt(10)
        p2.font.color.rgb = theme.text_secondary
        p2.font.name = theme.body_font
        p2.space_before = Pt(8)

    # ==================== SLIDE 4: ARSITEKTUR SISTEM ====================
    s4 = prs.slides.add_slide(blank_layout)
    add_theme_header(s4, theme, "Arsitektur Sistem: 5 Layer Pipeline End-to-End", "03 / SYSTEM ARCHITECTURE (PRD §3 & §5)")

    layers = [
        ("1. LAYER TELEMETRI HETEROGEN", "Web Servers (Nginx), SSH Auth (Syslog RFC 3164), Access Gate (CSV), DB Audit (PGAudit JSONL), & Hardware OT (MQTT Broker). 27 berkas log riil di sample_logs/.", theme.color_info),
        ("2. UNIVERSAL LOG ADAPTER", "app/log_adapter.py: Regex parser multi-format mengonversi log mentah ke kontrak kanonikal LogEvent secara dinamis relatif terhadap waktu nyata NOW().", theme.accent_primary),
        ("3. GATEKEEPER & VALIDATION", "app/main.py: Pydantic V2 schema validation, IPv4 format check, Time Drift verification (<24h). SOAR check: isolated host & suspended user -> HTTP 403 Forbidden.", theme.color_danger),
        ("4. THE BRAIN: CORRELATION ENGINE", "app/engine.py: Stateful in-memory sliding windows (10s, 60s, 1h), debouncing alert cooldown, Stateful Watchlist APT, dan korelasi sabotase OT.", theme.accent_secondary),
        ("5. STORAGE & VISUALISASI SOC", "Hot In-Memory Cache (5000 logs, 1000 alerts) + Cold Storage Archive. Dashboard web cyberpunk real-time (polling 2 detik) + SOAR Active Defense Panel.", theme.color_success)
    ]

    for idx, (ltitle, ldesc, lcolor) in enumerate(layers):
        y = Inches(1.6 + idx * 1.02)
        add_card_box(s4, theme, Inches(0.8), y, Inches(11.7), Inches(0.88), side_stripe_color=lcolor)

        tb = s4.shapes.add_textbox(Inches(1.1), y + Inches(0.1), Inches(11.2), Inches(0.68))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = ltitle
        p.font.size = Pt(11.5)
        p.font.bold = True
        p.font.color.rgb = theme.text_primary
        p.font.name = theme.title_font

        p2 = tf.add_paragraph()
        p2.text = ldesc
        p2.font.size = Pt(10)
        p2.font.color.rgb = theme.text_secondary
        p2.font.name = theme.body_font
        p2.space_before = Pt(3)

    # ==================== SLIDE 5: SPESIFIKASI DATA & GATEKEEPER ====================
    s5 = prs.slides.add_slide(blank_layout)
    add_theme_header(s5, theme, "Spesifikasi Kontrak Data Kanonikal & Ketahanan Gatekeeper", "04 / DATA CONTRACTS & VALIDATION (PRD §2 & §3A)")

    # Left: Canonical JSON Contract
    add_card_box(s5, theme, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.0), bg_color=theme.card_bg, border_color=theme.accent_primary)
    tb_l = s5.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(5.2), Inches(4.6))
    tf_l = tb_l.text_frame
    tf_l.word_wrap = True
    pl = tf_l.paragraphs[0]
    pl.text = "Kontrak Data Kanonikal (JSON Schema)"
    pl.font.size = Pt(13)
    pl.font.bold = True
    pl.font.color.rgb = theme.accent_primary
    pl.font.name = theme.title_font

    code_lines = [
        "{\n  \"timestamp\": \"YYYY-MM-DDThh:mm:ssZ\",",
        "  \"event_type\": \"login_attempt | web_request |",
        "                database_download | recon_probe |",
        "                ot_command | mqtt_command\",",
        "  \"user_id\": \"string (non-empty)\",",
        "  \"ip_address\": \"string (IPv4 Valid)\",",
        "  \"location\": \"string (non-empty)\",",
        "  \"status\": \"success | failed\",",
        "  // Opsional untuk Telemetri Hardware OT:",
        "  \"protocol\": \"MQTT | Modbus | BACnet\",",
        "  \"target_device\": \"CRAC-AC-CENTRAL-01\",",
        "  \"command\": \"POWER_OFF | STATUS_REPORT\"",
        "}"
    ]
    p_code = tf_l.add_paragraph()
    p_code.text = "\n".join(code_lines)
    p_code.font.size = Pt(9.5)
    p_code.font.color.rgb = theme.text_primary
    p_code.font.name = theme.code_font
    p_code.space_before = Pt(8)

    # Right: Gatekeeper Safeguards
    add_card_box(s5, theme, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.0), bg_color=theme.card_bg, border_color=theme.accent_secondary)
    tb_r = s5.shapes.add_textbox(Inches(7.0), Inches(1.8), Inches(5.3), Inches(4.6))
    tf_r = tb_r.text_frame
    tf_r.word_wrap = True
    pr = tf_r.paragraphs[0]
    pr.text = "Ketahanan Gatekeeper (Validasi Negatif)"
    pr.font.size = Pt(13)
    pr.font.bold = True
    pr.font.color.rgb = theme.accent_secondary
    pr.font.name = theme.title_font

    checks = [
        ("Format IP IPv4 Strict", "Menolak alamat string rusak atau bukan format desimal 4 oktet (misal '999.999.999.999')."),
        ("Missing / Null Fields", "Menolak paket jika atribut wajib seperti ip_address, user_id, atau location bernilai null atau kosong."),
        ("Time Drift Error (>24 Jam)", "Mencegah injeksi log basi lampau (seperti baterai RTC router CMOS reset ke 1970) dengan HTTP 422."),
        ("Future Drift Clock Error", "Menolak log yang memiliki stempel waktu di masa depan (>30 detik dari waktu sistem riil)."),
        ("SOAR Active Enforcement", "Jika IP terdaftar di isolated_hosts atau user_id di suspended_users, paket langsung DITOLAK di gerbang dengan HTTP 403 Forbidden.")
    ]

    for title, desc in checks:
        p_c = tf_r.add_paragraph()
        p_c.text = f"• {title}: {desc}"
        p_c.font.size = Pt(9.5)
        p_c.font.color.rgb = theme.text_secondary
        p_c.font.name = theme.body_font
        p_c.space_before = Pt(7)

    # ==================== SLIDE 6: 7 ATURAN KORELASI ====================
    s6 = prs.slides.add_slide(blank_layout)
    add_theme_header(s6, theme, "Matriks 7 Aturan Inti Mesin Korelasi (The Brain)", "05 / CORRELATION ENGINE (PRD §3C)")

    rows = 8
    cols = 4
    table_shape = s6.shapes.add_table(rows, cols, Inches(0.8), Inches(1.6), Inches(11.7), Inches(5.0))
    table = table_shape.table
    table.columns[0].width = Inches(1.8)
    table.columns[1].width = Inches(2.8)
    table.columns[2].width = Inches(1.8)
    table.columns[3].width = Inches(5.3)

    headers = ["ID Aturan", "Nama Insiden", "Severity", "Kondisi / Ambang Batas Korelasi"]
    for idx, h in enumerate(headers):
        cell = table.cell(0, idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = theme.card_alt
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = theme.accent_primary
        p.font.name = theme.title_font

    rule_data = [
        ("Rule 1", "DDoS / Botnet Flood", "TINGGI", "> 10 web_request dari 1 IP dalam rentang 10 detik."),
        ("Rule 2", "Impossible Travel", "KRITIS", "2 login_attempt sukses user sama di lokasi beda selang < 1 jam."),
        ("Rule 3", "Distributed Brute Force", "TINGGI", "> 5 login gagal untuk 1 user dari > 3 IP berbeda dalam 1 menit."),
        ("Rule 4", "Off-Hours Anomaly", "MENENGAH", "Aktivitas sukses user manusia di luar jam kerja (00:00 - 05:00)."),
        ("Rule 5", "Data Exfiltration", "KRITIS", "1 user_id download database > 50 kali dalam 1 menit."),
        ("APT Watchlist", "Low & Slow APT Correlation", "KRITIS", "Korelasi recon di Cold Storage dengan exploit di Hot Storage via Greylist."),
        ("Rule OT Sabotage", "Datacenter AC Shutdown", "KRITIS", "Perintah POWER_OFF / SHUTDOWN hardware pendingin via protokol MQTT.")
    ]

    for row_idx, r in enumerate(rule_data):
        for col_idx, text in enumerate(r):
            cell = table.cell(row_idx + 1, col_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = theme.card_bg
            p = cell.text_frame.paragraphs[0]
            p.text = text
            p.font.size = Pt(9.5)
            p.font.name = theme.body_font
            if col_idx == 2:
                p.font.bold = True
                if text == "KRITIS":
                    p.font.color.rgb = theme.color_danger
                elif text == "TINGGI":
                    p.font.color.rgb = theme.color_warning
                else:
                    p.font.color.rgb = theme.color_info
            else:
                p.font.color.rgb = theme.text_primary if col_idx < 2 else theme.text_secondary

    # ==================== SLIDE 7: SKENARIO 1 s/d 3 ====================
    s7 = prs.slides.add_slide(blank_layout)
    add_theme_header(s7, theme, "Bedah Skenario 1 s/d 3: Trafik Normal, DDoS Attack & Impossible Travel", "06 / DEMO SCENARIOS (PRD §4)")

    sc13 = [
        ("SKENARIO 1: TRAFIK NORMAL", "Baseline Operasional Sah",
         "• Sumber Data: sample_logs/web_access_normal_[1..3].log (Nginx)\n• Aktivitas: Karyawan mengakses portal HR, e-commerce, & customer desk pada jam kerja wajar.\n• Target Verifikasi: Membuktikan SIEM tidak menghasilkan alarm palsu (False Positive Zero).\n• Status SIEM: 0 Alarm Dipicu (Trafik Lolos Normal).",
         theme.color_success),
        ("SKENARIO 2: DDOS / BOTNET FLOOD", "Memicu Rule 1 (Severity: TINGGI)",
         "• Sumber Data: sample_logs/nginx_ddos_flood_[1..3].log (HTTP Flood)\n• Serangan: IP penyerang (45.33.32.156) membombardir endpoint API > 10 request dalam 10 detik.\n• Respon SIEM: Memicu alarm 'DDoS / Botnet' TINGGI.\n• Solusi SOC: Memicu Playbook SOAR Isolate Host untuk memutus trafik penyerang.",
         theme.color_warning),
        ("SKENARIO 3: IMPOSSIBLE TRAVEL", "Memicu Rule 2 (Severity: KRITIS)",
         "• Sumber Data: sample_logs/vpn_and_gate_access_[1..3].csv (RFID & VPN)\n• Anomali: User 'sarah_security' login fisik di Jakarta ID, lalu 1 detik kemudian login VPN dari Frankfurt DE.\n• Kecepatan Fisik: Mustahil berpindah ribuan km dalam 1 detik (pembajakan kredensial).\n• Respon SIEM: Memicu alarm Kritis & Playbook Suspend User.",
         theme.color_danger)
    ]

    for idx, (stitle, ssub, sbody, scolor) in enumerate(sc13):
        x = Inches(0.8 + idx * 3.95)
        add_card_box(s7, theme, x, Inches(1.6), Inches(3.75), Inches(5.0), side_stripe_color=scolor)

        tb = s7.shapes.add_textbox(x + Inches(0.25), Inches(1.8), Inches(3.3), Inches(4.6))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = stitle
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = theme.text_primary
        p.font.name = theme.title_font

        p_sub = tf.add_paragraph()
        p_sub.text = ssub
        p_sub.font.size = Pt(9.5)
        p_sub.font.italic = True
        p_sub.font.color.rgb = scolor
        p_sub.font.name = theme.body_font

        p_desc = tf.add_paragraph()
        p_desc.text = sbody
        p_desc.font.size = Pt(9.5)
        p_desc.font.color.rgb = theme.text_secondary
        p_desc.font.name = theme.body_font
        p_desc.space_before = Pt(10)

    # ==================== SLIDE 8: SKENARIO 4 s/d 6 ====================
    s8 = prs.slides.add_slide(blank_layout)
    add_theme_header(s8, theme, "Bedah Skenario 4 s/d 6: Brute Force, Exfiltration & Careless Whitelisting", "07 / DEMO SCENARIOS (PRD §4)")

    sc46 = [
        ("SKENARIO 4: DISTRIBUTED BRUTE FORCE", "Memicu Rule 3 (Severity: TINGGI)",
         "• Sumber Data: sample_logs/auth_sshd_bruteforce_[1..3].log (Syslog RFC 3164)\n• Taktik Penyerang: Melancarkan serangan tebak password SSH terdistribusi melalui 4 node Tor / VPS.\n• Kriteria SIEM: > 5 kegagalan login untuk 1 user dari > 3 IP berbeda dalam rentang 60 detik.\n• Hasil: Alarm 'Distributed Brute Force' berhasil dibunyikan.",
         theme.color_warning),
        ("SKENARIO 5: DATA EXFILTRATION", "Memicu Rule 5 (Severity: KRITIS)",
         "• Sumber Data: sample_logs/postgres_audit_exfiltration_[1..3].jsonl (PGAudit)\n• Ancaman Insider: Oknum developer (dev_rogue) mendownload data PII nasabah > 50 kali dalam 1 menit.\n• Respon SIEM: Deteksi lonjakan volume unduhan data sensitif dan membunyikan alarm KRITIS.\n• Mitigasi: Pembatasan akses database & suspend kredensial user.",
         theme.color_danger),
        ("SKENARIO 6: CARELESS WHITELISTING", "Demonstrasi Fatal Error & Blind Spot",
         "• Konsep Edukasi: Admin memasukkan IP penyerang ke daftar Whitelist (/api/whitelist).\n• Dampak: Korelasi engine otomatis MENGABAIKAN log dari IP tersebut (Zero Alert).\n• Pembuktian SOC: Menunjukkan risiko fatal whitelisting ceroboh tanpa review keamanan bertingkat.\n• Pemulihan: Un-whitelist mengembalikan alarm.",
         theme.accent_secondary)
    ]

    for idx, (stitle, ssub, sbody, scolor) in enumerate(sc46):
        x = Inches(0.8 + idx * 3.95)
        add_card_box(s8, theme, x, Inches(1.6), Inches(3.75), Inches(5.0), side_stripe_color=scolor)

        tb = s8.shapes.add_textbox(x + Inches(0.25), Inches(1.8), Inches(3.3), Inches(4.6))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = stitle
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = theme.text_primary
        p.font.name = theme.title_font

        p_sub = tf.add_paragraph()
        p_sub.text = ssub
        p_sub.font.size = Pt(9.5)
        p_sub.font.italic = True
        p_sub.font.color.rgb = scolor
        p_sub.font.name = theme.body_font

        p_desc = tf.add_paragraph()
        p_desc.text = sbody
        p_desc.font.size = Pt(9.5)
        p_desc.font.color.rgb = theme.text_secondary
        p_desc.font.name = theme.body_font
        p_desc.space_before = Pt(10)

    # ==================== SLIDE 9: SKENARIO 7 & 8 ====================
    s9 = prs.slides.add_slide(blank_layout)
    add_theme_header(s9, theme, "Bedah Skenario 7 & 8: Data Cacat & The Low and Slow APT", "08 / DEMO SCENARIOS (PRD §4)")

    # Left: Scenario 7
    add_card_box(s9, theme, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.0), side_stripe_color=theme.color_warning)
    tb7 = s9.shapes.add_textbox(Inches(1.05), Inches(1.8), Inches(5.15), Inches(4.6))
    tf7 = tb7.text_frame
    tf7.word_wrap = True
    p7 = tf7.paragraphs[0]
    p7.text = "SKENARIO 7: DATA CACAT & TIME DRIFT"
    p7.font.size = Pt(12.5)
    p7.font.bold = True
    p7.font.color.rgb = theme.text_primary
    p7.font.name = theme.title_font

    sub7 = tf7.add_paragraph()
    sub7.text = "Uji Ketahanan Gatekeeper terhadap Log Korup (HTTP 422)"
    sub7.font.size = Pt(10)
    sub7.font.italic = True
    sub7.font.color.rgb = theme.color_warning
    sub7.font.name = theme.body_font

    p7_body = tf7.add_paragraph()
    p7_body.text = (
        "• Sumber Data: sample_logs/legacy_corrupted_device_[1..3].log\n\n"
        "• Kasus Uji 1 (Malformed JSON): Data terpotong di tengah transmisi jaringan.\n"
        "• Kasus Uji 2 (Missing IP): Log telemetri tanpa atribut alamat IP pengirim.\n"
        "• Kasus Uji 3 (Invalid IP format): Format alamat IP bukan IPv4 desimal valid.\n"
        "• Kasus Uji 4 (CMOS Battery Reset): Timestamp perangkat lawas reset ke 1970 (> 24 jam lampau).\n"
        "• Kasus Uji 5 (Future Time Drift): Jam perangkat melompat ke masa depan.\n\n"
        "✔ Hasil Gatekeeper: 100% data cacat ditolak dengan respons HTTP 422 tanpa menyebabkan server crash."
    )
    p7_body.font.size = Pt(9.5)
    p7_body.font.color.rgb = theme.text_secondary
    p7_body.font.name = theme.body_font
    p7_body.space_before = Pt(8)

    # Right: Scenario 8
    add_card_box(s9, theme, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.0), side_stripe_color=theme.accent_secondary)
    tb8 = s9.shapes.add_textbox(Inches(7.05), Inches(1.8), Inches(5.25), Inches(4.6))
    tf8 = tb8.text_frame
    tf8.word_wrap = True
    p8 = tf8.paragraphs[0]
    p8.text = "SKENARIO 8: THE LOW & SLOW APT"
    p8.font.size = Pt(12.5)
    p8.font.bold = True
    p8.font.color.rgb = theme.text_primary
    p8.font.name = theme.title_font

    sub8 = tf8.add_paragraph()
    sub8.text = "Jebakan Hot vs Cold Storage & Solusi Stateful Watchlist"
    sub8.font.size = Pt(10)
    sub8.font.italic = True
    sub8.font.color.rgb = theme.accent_secondary
    sub8.font.name = theme.body_font

    p8_body = tf8.add_paragraph()
    p8_body.text = (
        "• Masalah Arsitektur: Log hanya disimpan di Hot Storage selama 7-14 hari demi efisiensi biaya. Serangan APT berjarak 20 hari lolos karena data awal sudah diarsipkan ke Cold Storage (Temporal Blind Spot).\n\n"
        "• Alur Serangan Dua Babak:\n"
        "  - Tanggal 1: Penyerang melakukan Reconnaissance Probe (/api/v1/recon).\n"
        "  - Tanggal 20: Penyerang meluncurkan Exploit Request.\n\n"
        "✔ Solusi Mini-SIEM: Begitu recon terdeteksi pada Tanggal 1, IP otomatis dicatat ke Stateful Watchlist (Greylist). Saat exploit tiba di Tanggal 20, alarm APT Kritis dipicu dan riwayat dingin direhidrasi untuk pembuktian forensik."
    )
    p8_body.font.size = Pt(9.5)
    p8_body.font.color.rgb = theme.text_secondary
    p8_body.font.name = theme.body_font
    p8_body.space_before = Pt(8)

    # ==================== SLIDE 10: SKENARIO 9 (OT SABOTAGE AC CENTRAL) ====================
    s10 = prs.slides.add_slide(blank_layout)
    add_theme_header(s10, theme, "Bedah Skenario 9: OT Hardware Sabotage - AC Central Shutdown (MQTT)", "09 / DEMO SCENARIOS (PRD §4)")

    sc9_cols = [
        ("1. KONTEKS HARDWARE OT", theme.color_info,
         "• Target Fisik: Unit pendingin utama datacenter (CRAC-AC-CENTRAL-01, Chiller Plant, atau PAC In-Row Cooling).\n\n"
         "• Protokol Komunikasi: MQTT (Message Queuing Telemetry Transport) via broker terhubung ke sensor BMS/SCADA.\n\n"
         "• Dampak Kegagalan: Jika AC mati mendadak, suhu rak server melonjak drastis memicu thermal runaway (server overheat, mati massal, & kerusakan hardware bernilai miliaran)."),
        ("2. ALUR REPLAY 2 TAHAP", theme.color_warning,
         "• Pola mirip Skenario 3 (Akses Fisik vs Remote):\n\n"
         "• Event 1 (Status Monitoring):\n"
         "  Sensor hardware membaca status telemetri suhu (18.5°C) & perintah STATUS_REPORT via topik MQTT datacenter/cooling/crac01/telemetry -> Diterima normal (False Positive Zero).\n\n"
         "• Event 2 (Perintah Sabotase):\n"
         "  Oknum ilegal mengirim perintah POWER_OFF via topik datacenter/cooling/crac01/power."),
        ("3. DETEKSI & PERTAHANAN SOAR", theme.color_danger,
         "• Deteksi Mesin SIEM:\n"
         "  SIEM mengenali protokol MQTT dan jenis hardware pendingin OT, langsung memicu alarm KRITIS: 'OT Sabotage: Datacenter AC Central Shutdown (MQTT)'.\n\n"
         "• Respons Aktif SOAR:\n"
         "  Analis mengisolasi IP penyerang (192.168.30.99) via SOAR Host Quarantine. Perintah MQTT sabotase lanjutan langsung di-drop di Gatekeeper (HTTP 403)!")
    ]

    for idx, (col_title, col_color, col_body) in enumerate(sc9_cols):
        x = Inches(0.8 + idx * 3.95)
        add_card_box(s10, theme, x, Inches(1.6), Inches(3.75), Inches(5.0), side_stripe_color=col_color)

        tb = s10.shapes.add_textbox(x + Inches(0.25), Inches(1.8), Inches(3.3), Inches(4.6))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = col_title
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = theme.text_primary
        p.font.name = theme.title_font

        p_desc = tf.add_paragraph()
        p_desc.text = col_body
        p_desc.font.size = Pt(9.5)
        p_desc.font.color.rgb = theme.text_secondary
        p_desc.font.name = theme.body_font
        p_desc.space_before = Pt(10)

    # ==================== SLIDE 11: SOAR ACTIVE DEFENSE ====================
    s11 = prs.slides.add_slide(blank_layout)
    add_theme_header(s11, theme, "SOAR Active Defense: Network Quarantine & Identity Containment", "10 / INCIDENT RESPONSE PLAYBOOKS (PRD §3D)")

    soar_items = [
        ("1. Network Quarantine (Isolate Host)", "POST /api/soar/isolate",
         "Analis mengisolasi IP penyerang langsung dari kartu alert dashboard atau API. Status alert bertransisi menjadi [CONTAINED] dan buffer korelasi aktif dibersihkan untuk mencegah ghost alerts. Seluruh log lanjutan dari IP tersebut ditolak otomatis di Gatekeeper (HTTP 403 HOST_ISOLATED).", theme.color_danger),
        ("2. Identity Containment (Suspend User)", "POST /api/soar/suspend-user",
         "Untuk insiden kredensial terkompromi (seperti Impossible Travel), akun user dibekukan seketika. Setiap percobaan autentikasi atau aktivitas lanjutan yang mengatasnamakan user tersebut ditolak di gerbang Gatekeeper (HTTP 403 USER_SUSPENDED).", theme.accent_secondary),
        ("3. Access Restoration (Release Playbook)", "POST /api/soar/release & unsuspend",
         "Setelah ancaman dinetralisir dan investigasi forensik selesai, analis dapat melepas status karantina host atau mengaktifkan kembali akun pengguna dengan satu klik tanpa perlu merestart sistem.", theme.color_success),
        ("4. Immutable Audit Trail", "GET /api/soar/audit",
         "Setiap aksi karantina, pembekuan user, dan pemulihan dicatat secara kronologis dalam jejak audit SOAR (mencakup action, target, alasan analis, dan stempel waktu ISO UTC) untuk akuntabilitas operasional.", theme.accent_primary)
    ]

    for idx, (title, endpoint, desc, color) in enumerate(soar_items):
        y = Inches(1.6 + idx * 1.25)
        add_card_box(s11, theme, Inches(0.8), y, Inches(11.7), Inches(1.1), side_stripe_color=color)

        tb = s11.shapes.add_textbox(Inches(1.1), y + Inches(0.12), Inches(11.2), Inches(0.9))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = f"{title} — [{endpoint}]"
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = theme.text_primary
        p.font.name = theme.title_font

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(10)
        p2.font.color.rgb = theme.text_secondary
        p2.font.name = theme.body_font
        p2.space_before = Pt(4)

    # ==================== SLIDE 12: TATA KELOLA AI & HUMAN-IN-THE-LOOP ====================
    s12 = prs.slides.add_slide(blank_layout)
    add_theme_header(s12, theme, "Tata Kelola & Arsitektur Pengaman: Human-Supervised AI Adaptation", "11 / GOVERNANCE & AI SAFETY (PRD §6)")

    risks = [
        ("1. Concept Drift Poisoning", "Penyerang APT menyuntikkan anomali berfrekuensi rendah bertahap (10 -> 500 file/minggu). AI tanpa supervisi menganggap tren ini normal.", theme.color_danger),
        ("2. Self-Inflicted DoS", "Aturan adaptif salah mengklasifikasikan lonjakan checkout promo diskon (Flash Sale) sebagai DDoS, melumpuhkan transaksi pelanggan sah.", theme.color_warning),
        ("3. Pelanggaran Audit Kepatuhan", "ISO 27001 & PCI-DSS melarang keputusan 'Black-Box' tanpa tiket persetujuan formal dan verifikator penanggung jawab yang jelas.", theme.accent_secondary)
    ]

    for idx, (rtitle, rdesc, rcolor) in enumerate(risks):
        x = Inches(0.8 + idx * 3.95)
        add_card_box(s12, theme, x, Inches(1.6), Inches(3.75), Inches(2.2), side_stripe_color=rcolor)

        tb = s12.shapes.add_textbox(x + Inches(0.25), Inches(1.75), Inches(3.3), Inches(1.9))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = rtitle
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = theme.text_primary
        p.font.name = theme.title_font

        p_d = tf.add_paragraph()
        p_d.text = rdesc
        p_d.font.size = Pt(9.5)
        p_d.font.color.rgb = theme.text_secondary
        p_d.font.name = theme.body_font
        p_d.space_before = Pt(6)

    # Blueprint 5-Step Cycle Box
    add_card_box(s12, theme, Inches(0.8), Inches(4.0), Inches(11.7), Inches(2.7), border_color=theme.accent_primary)
    tb_c = s12.shapes.add_textbox(Inches(1.0), Inches(4.15), Inches(11.3), Inches(2.4))
    tf_c = tb_c.text_frame
    tf_c.word_wrap = True
    p = tf_c.paragraphs[0]
    p.text = "🏛️ Blueprint Siklus Aman: Dari Deteksi Kasus Baru hingga Implementasi Produksi"
    p.font.size = Pt(12.5)
    p.font.bold = True
    p.font.color.rgb = theme.accent_primary
    p.font.name = theme.title_font

    steps_text = (
        "1. Gatekeeper Ingest: Validasi skema deterministik (Tolak HTTP 422 jika data rusak)\n"
        "2. Dead-Letter Queue (DLQ): Payload anomali baru diisolasi ke antrean karantina terpisah tanpa mengganggu pipeline utama\n"
        "3. AI Suggestion Engine: Sintesis pola klaster anomali dan perumusan 'Candidate Rule' secara terprogram\n"
        "4. Canary Replay Test: Pengujian regresi aturan pada 30 hari data historis; auto-rollback seketika jika rasio False Positive > 2%\n"
        "5. HITL Approval (Shadow Mode): Aturan beroperasi dalam mode bayangan hingga Lead Analis SOC mengotorisasi promosi ke produksi"
    )
    p_steps = tf_c.add_paragraph()
    p_steps.text = steps_text
    p_steps.font.size = Pt(10)
    p_steps.font.color.rgb = theme.text_secondary
    p_steps.font.name = theme.body_font
    p_steps.space_before = Pt(6)

    # ==================== SLIDE 13: PRASYARAT INTEGRASI INFRASTRUKTUR ====================
    s13 = prs.slides.add_slide(blank_layout)
    add_theme_header(s13, theme, "Prasyarat Integrasi Infrastruktur: Universal Ingestion Syllabus", "12 / INTEGRATION SYLLABUS (PRD §7)")

    syllabus = [
        ("Tahap 1: Fondasi Bahan Baku Log", theme.accent_primary,
         "• Anatomi Log: Membedakan Event (kejadian), Payload (isi utama), dan Metadata (keterangan tambahan).\n• Log Lifecycle: Titik lahir -> Forwarding -> Queueing -> Hot Storage -> Cold Archive.\n• Kalkulasi EPS (Events Per Second): Menghitung beban volume masuk agar mesin korelasi tidak crash."),
        ("Tahap 2: Standarisasi Bahasa", theme.color_info,
         "• Protokol RFC 5424: Format pembungkus Syslog terstandarisasi (Timestamp, Hostname, Severity, Payload).\n• Kepatuhan NIST SP 800-92 & ISO 27002: Wajib catat kegagalan login & eskalasi privilese.\n• OCSF / CIM Universal Schema: Menyatukan field multi-vendor (SrcIP & Client_Address -> src_endpoint.ip)."),
        ("Tahap 3: Prasyarat Infrastruktur Jaringan", theme.accent_secondary,
         "• Sinkronisasi Waktu Sentral (NTP): Syarat mutlak! Seluruh server wajib merujuk waktu atomik UTC yang sama. Deviasi 1s merusak urutan korelasi.\n• Log Forwarder Agregator: Menggunakan Logstash/Fluentd terdistribusi di tiap cabang.\n• Noise Filtering: Membuang ping rutin internal di level forwarder."),
        ("Tahap 4: Prasyarat Bisnis & Konteks", theme.color_success,
         "• Asset Management: Klasifikasi kekritisan server inti vs WiFi publik untuk severity proporsional.\n• Definisi Use Case: Merumuskan skenario deteksi ancaman sebelum sistem dioperasikan.\n• SOP & Eskalasi Bertingkat: Penetapan tim penanggung jawab (Jaringan, Aplikasi, SOC) saat alert kritis terpicu.")
    ]

    for idx, (stitle, scolor, sdesc) in enumerate(syllabus):
        col = idx % 2
        row = idx // 2
        x = Inches(0.8 + col * 5.95)
        y = Inches(1.6 + row * 2.5)

        add_card_box(s13, theme, x, y, Inches(5.75), Inches(2.35), side_stripe_color=scolor)

        tb = s13.shapes.add_textbox(x + Inches(0.28), y + Inches(0.15), Inches(5.3), Inches(2.05))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = stitle
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = theme.text_primary
        p.font.name = theme.title_font

        p_desc = tf.add_paragraph()
        p_desc.text = sdesc
        p_desc.font.size = Pt(9.5)
        p_desc.font.color.rgb = theme.text_secondary
        p_desc.font.name = theme.body_font
        p_desc.space_before = Pt(6)

    # ==================== SLIDE 14: STORAGE BERTINGKAT (HOT VS COLD) ====================
    s14 = prs.slides.add_slide(blank_layout)
    add_theme_header(s14, theme, "Arsitektur Storage Bertingkat: Hot Storage vs Cold Storage", "13 / STORAGE ARCHITECTURE (PRD §8)")

    # Left: The Brute-Force Fallacy
    add_card_box(s14, theme, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.0), side_stripe_color=theme.color_danger)
    tb_bf = s14.shapes.add_textbox(Inches(1.05), Inches(1.8), Inches(5.15), Inches(4.6))
    tf_bf = tb_bf.text_frame
    tf_bf.word_wrap = True
    p = tf_bf.paragraphs[0]
    p.text = "⚠️ Limitasi Skalabilitas Linier (The Brute-Force Fallacy)"
    p.font.size = Pt(12.5)
    p.font.bold = True
    p.font.color.rgb = theme.text_primary
    p.font.name = theme.title_font

    bf_text = (
        "Asumsi umum: 'Jika peretas menyerang di Hari ke-20, perbesar retensi Hot Storage menjadi 30 atau 90 hari.'\n\n"
        "Mengapa pendekatan ini gagal di skala perusahaan?\n\n"
        "1. Biaya Eksponensial: Menyimpan petabyte log mentah di tier RAM/NVMe SSD menaikkan biaya infrastruktur data secara tidak proporsional.\n\n"
        "2. Degradasi Latensi Kueri: Indeks pencarian yang membengkak menurunkan performa evaluasi aturan korelasi (dari < 10ms menjadi detik/menit).\n\n"
        "3. Rasio Noise-to-Signal Tinggi: 98% data adalah telemetri rutin; mengalokasikan tier premium untuk data non-kritis memboroskan anggaran."
    )
    p_bf = tf_bf.add_paragraph()
    p_bf.text = bf_text
    p_bf.font.size = Pt(9.5)
    p_bf.font.color.rgb = theme.text_secondary
    p_bf.font.name = theme.body_font
    p_bf.space_before = Pt(8)

    # Right: Dual Solutions
    add_card_box(s14, theme, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.0), side_stripe_color=theme.color_success)
    tb_sol = s14.shapes.add_textbox(Inches(7.05), Inches(1.8), Inches(5.25), Inches(4.6))
    tf_sol = tb_sol.text_frame
    tf_sol.word_wrap = True
    p = tf_sol.paragraphs[0]
    p.text = "✔ Dua Solusi Arsitektural Mini-SIEM"
    p.font.size = Pt(12.5)
    p.font.bold = True
    p.font.color.rgb = theme.text_primary
    p.font.name = theme.title_font

    sol_text = (
        "1. Stateful Entity Tracking (Lightweight Memory):\n"
        "• Alih-alih menahan puluhan gigabyte log mentah, SIEM hanya menyimpan satu baris JSON status di tabel persisten:\n"
        "  {\"ip\": \"198.51.100.77\", \"risk\": \"HIGH\", \"score\": 85, \"reason\": \"recon_day_1\"}\n"
        "• Jejak memori minimal ini dipertahankan lintas siklus retensi untuk dicocokkan saat entitas aktif kembali di Hari ke-20.\n\n"
        "2. Cold Log Rehydration (On-Demand Forensic):\n"
        "• Saat ancaman terverifikasi di Hari ke-20, sistem secara terarah mengekstraksi paket log historis dari arsip dingin:\n"
        "  cold_s3://logs/{yyyy}/{mm}/{dd}/{ip_hash}/\n"
        "• Log ditarik kembali ke memori aktif analis untuk rekonstruksi linimasa serangan (Kill Chain) tanpa full-table scan."
    )
    p_sol = tf_sol.add_paragraph()
    p_sol.text = sol_text
    p_sol.font.size = Pt(9.5)
    p_sol.font.color.rgb = theme.text_secondary
    p_sol.font.name = theme.body_font
    p_sol.space_before = Pt(8)

    # ==================== SLIDE 15: ROADMAP PM 13 FITUR ====================
    s15 = prs.slides.add_slide(blank_layout)
    add_theme_header(s15, theme, "Roadmap Ekosistem SIEM & Matriks PM 13 Fitur", "14 / PRODUCT ROADMAP (PRD §9)")

    phases = [
        ("MVP v1: Foundation (Fitur 1 - 6)", theme.accent_primary,
         "• 1. Log Ingestion Engine (Critical)\n• 2. Data Parsing & Normalization (Critical)\n• 3. Hot Storage & Indexing (Critical)\n• 4. Search & Query Interface (High)\n• 5. Rule-Based Alerting (High)\n• 6. Basic Dashboard & Reporting (Medium)"),
        ("Phase 2: Intelligence (Fitur 7 - 10)", theme.accent_secondary,
         "• 7. Correlation Engine (High - Stateful)\n• 8. Threat Intel Integration (High - IoC)\n• 9. Incident & Case Management (Medium)\n• 10. Cold Storage / Archiving (Medium - Retensi 1-7 Tahun)"),
        ("Phase 3: Advanced (Fitur 11 - 13)", theme.color_success,
         "• 11. UEBA (High - Behavioral ML Model)\n• 12. SOAR (High - Automated Playbook)\n• 13. AI Adaptive Rule Generator (High - DLQ, Canary Replay, & HITL Approval)")
    ]

    for idx, (pname, pcolor, plist) in enumerate(phases):
        x = Inches(0.8 + idx * 3.95)
        add_card_box(s15, theme, x, Inches(1.6), Inches(3.75), Inches(2.7), side_stripe_color=pcolor)

        tb = s15.shapes.add_textbox(x + Inches(0.25), Inches(1.75), Inches(3.3), Inches(2.4))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = pname
        p.font.size = Pt(11.5)
        p.font.bold = True
        p.font.color.rgb = theme.text_primary
        p.font.name = theme.title_font

        p_l = tf.add_paragraph()
        p_l.text = plist
        p_l.font.size = Pt(9.5)
        p_l.font.color.rgb = theme.text_secondary
        p_l.font.name = theme.body_font
        p_l.space_before = Pt(6)

    # Bottom Dependency Logic
    add_card_box(s15, theme, Inches(0.8), Inches(4.5), Inches(11.7), Inches(2.2), border_color=theme.accent_primary)
    tb_dep = s15.shapes.add_textbox(Inches(1.0), Inches(4.65), Inches(11.3), Inches(1.9))
    tf_dep = tb_dep.text_frame
    tf_dep.word_wrap = True
    p = tf_dep.paragraphs[0]
    p.text = "💡 Tiga Prinsip Rantai Dependensi Manajemen Produk (PM Logic)"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = theme.accent_primary
    p.font.name = theme.title_font

    dep_points = (
        "1. Fondasi Data Sebelum Analitik Lanjut: Ingestion, Normalisasi, dan Penyimpanan Hot Storage harus stabil sebelum mengaktifkan UEBA. Tanpa data historis terstandarisasi, model perilaku tidak memiliki baseline yang andal.\n\n"
        "2. Pengendalian Risiko Sebelum Otomasi Penuh: SOAR diimplementasikan pada Phase 3 setelah Case Management dan Threat Intel matang. Mitigasi otomatis tanpa akurasi tinggi berisiko memblokir trafik operasional yang sah.\n\n"
        "3. Keseimbangan Performa & Efisiensi Biaya: Pemisahan Hot Storage (Phase 1) dan Cold Archive (Phase 2) menjaga performa korelasi real-time tetap optimal sekaligus mengendalikan biaya penyimpanan jangka panjang."
    )
    p_pts = tf_dep.add_paragraph()
    p_pts.text = dep_points
    p_pts.font.size = Pt(9.5)
    p_pts.font.color.rgb = theme.text_secondary
    p_pts.font.name = theme.body_font
    p_pts.space_before = Pt(6)

    # ==================== SLIDE 16: PENGUJIAN OTOMATIS & KESIMPULAN ====================
    s16 = prs.slides.add_slide(blank_layout)
    add_theme_header(s16, theme, "Verifikasi Kualitas (29 Tests Passed) & Kesimpulan Strategis", "15 / QA & CONCLUSION (PRD §5)")

    # Big Badge
    add_card_box(s16, theme, Inches(0.8), Inches(1.6), Inches(11.7), Inches(1.1), border_color=theme.color_success)
    tb_b = s16.shapes.add_textbox(Inches(1.0), Inches(1.7), Inches(11.3), Inches(0.9))
    tf_b = tb_b.text_frame
    p_b = tf_b.paragraphs[0]
    p_b.text = "✔ PYTEST STATUS: 29 PASSED (100% SUCCESS RATE IN 2.06s)"
    p_b.font.size = Pt(15)
    p_b.font.bold = True
    p_b.font.color.rgb = theme.color_success
    p_b.font.name = theme.title_font

    p_b2 = tf_b.add_paragraph()
    p_b2.text = "Rangkaian pengujian otomatis end-to-end mencakup validasi skema negatif, aturan korelasi sliding window, siklus hidup SOAR, dan skenario OT MQTT."
    p_b2.font.size = Pt(10.5)
    p_b2.font.color.rgb = theme.text_secondary
    p_b2.font.name = theme.body_font
    p_b2.space_before = Pt(4)

    # 4 Conclusion Cards
    conclusions = [
        ("Keandalan Gatekeeper Terbukti", "Validasi ketat kontrak skema data dan audit time drift mencegah polusi data telemetri sebelum masuk ke mesin korelasi.", theme.accent_primary),
        ("Deteksi Multi-Vektor Lengkap", "Mencakup spektrum serangan siber modern: DDoS L7, anomali identitas fisik vs remote, eksfiltrasi database, ancaman APT lintas waktu, hingga sabotase fisik OT/IoT.", theme.color_info),
        ("SOAR Active Defense Menghentikan Serangan", "SOC tidak hanya pasif memantau grafik melainkan mampu melakukan Network Quarantine dan User Suspension secara real-time di level Gatekeeper.", theme.color_success),
        ("Rekomendasi Skalabilitas Produksi", "Migrasi in-memory stateful store ke Redis/Kafka untuk ingestion jutaan event per detik, integrasi LLM Agent untuk playbook auto-remediation, dan ekspansi protokol OT (Modbus, BACnet).", theme.accent_secondary)
    ]

    for idx, (ctitle, cdesc, ccolor) in enumerate(conclusions):
        y = Inches(2.9 + idx * 0.98)
        add_card_box(s16, theme, Inches(0.8), y, Inches(11.7), Inches(0.85), side_stripe_color=ccolor)

        tb = s16.shapes.add_textbox(Inches(1.1), y + Inches(0.08), Inches(11.2), Inches(0.68))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = f"{idx+1}. {ctitle}"
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = theme.text_primary
        p.font.name = theme.title_font

        p2 = tf.add_paragraph()
        p2.text = cdesc
        p2.font.size = Pt(9.5)
        p2.font.color.rgb = theme.text_secondary
        p2.font.name = theme.body_font
        p2.space_before = Pt(2)

    prs.save(theme.output_file)
    print(f"[OK] Generated: {theme.output_file} ({len(prs.slides)} slides) - {theme.theme_name}")

def main():
    print("================================================================================")
    print("Generating 5 Variant Themes of Master Presentation (16 Slides Each)...")
    print("================================================================================")
    themes = get_themes()
    for theme in themes:
        build_variant_deck(theme)
    print("\nAll 5 Theme Variants Generated Successfully!")

if __name__ == "__main__":
    main()
